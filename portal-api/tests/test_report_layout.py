import io

from app.services import report_layout, report_render


def _h(client, pnr):
    return {"Authorization": "Bearer " + client.post("/auth/dev-login", json={"personalnummer": pnr}).json()["access_token"]}


def test_normalize_drops_unknown_and_appends_missing():
    cfg = report_layout.normalize({
        "accent": "rot", "sections": [{"key": "nps", "enabled": False, "title": "Mein NPS"}, {"key": "boese", "enabled": True}, {"key": "nps"}],
    })
    keys = [s["key"] for s in cfg["sections"]]
    assert keys[0] == "nps" and "boese" not in keys and keys.count("nps") == 1
    assert set(keys) == set(report_layout.SECTIONS)
    assert cfg["accent"] == "#78ab78"  # ungueltige Farbe -> Standard
    assert cfg["sections"][0] == {"key": "nps", "enabled": False, "title": "Mein NPS"}


def test_html_respects_order_titles_columns_and_toggles():
    cfg = report_layout.normalize({
        "title": "Bericht für {leader}", "footer": "Streng geheim", "columns": {"vorrunde": False, "stddev": True},
        "sections": [{"key": "guide", "enabled": True, "title": "Lesehilfe"}, {"key": "dimensions", "enabled": True, "title": "Alle Themen"},
                     {"key": "nps", "enabled": False, "title": "NPS"}, {"key": "highlights", "enabled": False}],
    })
    html = report_render.render_html(report_render.SAMPLE, cfg)
    assert "Bericht für Alex Beispiel" in html and "Streng geheim" in html
    assert html.index("Lesehilfe") < html.index("Alle Themen")
    assert "Weiterempfehlung" not in html and "Stärken und Potenziale" not in html
    assert "<th>Vorrunde</th>" not in html and "<th>Std.abw.</th>" in html


def test_pptx_xlsx_csv_exports_follow_layout_and_privacy_toggle():
    from openpyxl import load_workbook
    from pptx import Presentation

    cfg = report_layout.normalize({})
    pptx = Presentation(io.BytesIO(report_render.build_pptx(report_render.SAMPLE, cfg)))
    assert len(pptx.slides) >= 5
    wb = load_workbook(io.BytesIO(report_render.build_xlsx(report_render.SAMPLE, cfg)))
    assert {"Überblick", "Themen", "Fragen", "NPS"} <= set(wb.sheetnames)
    assert "Freitexte" not in wb.sheetnames  # Standard: rohe Freitexte nicht im Export
    csv = report_render.build_csv(report_render.SAMPLE, cfg)
    assert "Kommunikation" in csv and "Offene und ehrliche" not in csv

    cfg2 = report_layout.normalize({"sections": [{"key": "freetext", "enabled": True}]})
    assert "Freitexte" in load_workbook(io.BytesIO(report_render.build_xlsx(report_render.SAMPLE, cfg2))).sheetnames
    assert "Offene und ehrliche" in report_render.build_csv(report_render.SAMPLE, cfg2)


def test_layout_api_roundtrip_preview_and_samples(client, seeded_users):
    h = _h(client, "T-ADMIN")
    cur = client.get("/report-layout", headers=h).json()
    assert len(cur["catalog"]) == len(report_layout.SECTIONS)
    cfg = cur["config"]
    cfg["title"] = "Mein Titel {round}"
    assert client.put("/report-layout", json={"config": cfg}, headers=h).json()["config"]["title"] == "Mein Titel {round}"
    assert client.get("/report-layout", headers=h).json()["config"]["title"] == "Mein Titel {round}"
    assert "Mein Titel Beispiel-Runde" in client.post("/report-layout/preview", json={"config": cfg}, headers=h).json()["html"]
    for fmt, magic in (("pptx", b"PK"), ("xlsx", b"PK")):
        r = client.post(f"/report-layout/sample?format={fmt}", json={"config": cfg}, headers=h)
        assert r.status_code == 200 and r.content[:2] == magic
    assert client.post("/report-layout/sample?format=csv", json={"config": cfg}, headers=h).text.startswith("﻿")
    assert client.post("/report-layout/sample?format=exe", json={"config": cfg}, headers=h).status_code == 400
    assert client.post("/report-layout/reset", headers=h).json()["config"]["title"] == report_layout.DEFAULT["title"]


def test_layout_needs_surveys_manage(client, seeded_users):
    h = _h(client, "T-ADMIN")
    rid = next(r["id"] for r in client.get("/access/roles", headers=h).json() if r["name"] == "Nur Auswertung")
    client.put(f"/access/persons/{seeded_users['employee'].id}", json={"role_ids": [rid]}, headers=h)
    hv = _h(client, "T-MA")
    assert client.get("/report-layout", headers=hv).status_code == 403
    assert client.put("/report-layout", json={"config": {}}, headers=hv).status_code == 403
    # Nicht-Admin
    assert client.get("/report-layout", headers=_h(client, "T-FK")).status_code == 403


def test_benchmark_export_formats_for_results_viewers(client, seeded_users, db_session):
    from datetime import datetime, timezone
    from app.models.round import Round, RoundStatus
    from app.models.survey import SurveyTemplate, SurveyVersion
    t = SurveyTemplate(name="T")
    db_session.add(t)
    db_session.flush()
    v = SurveyVersion(survey_template_id=t.id, version_number=1)
    db_session.add(v)
    db_session.flush()
    r = Round(name="R1", survey_version_id=v.id, status=RoundStatus.berichtet,
              start_at=datetime(2026, 1, 1, tzinfo=timezone.utc), end_at=datetime(2026, 2, 1, tzinfo=timezone.utc))
    db_session.add(r)
    db_session.commit()
    h = _h(client, "T-ADMIN")
    assert client.get(f"/benchmark/export.xlsx?round_id={r.id}", headers=h).content[:2] == b"PK"
    assert client.get(f"/benchmark/export.pptx?round_id={r.id}", headers=h).content[:2] == b"PK"
    assert client.get("/benchmark/export.xlsx?round_id=999", headers=h).status_code == 404


def test_custom_text_blocks_overrides_placeholders_and_escaping():
    cfg = report_layout.normalize({
        "texts": {"guide_body": "Hallo {leader}, dein bestes Thema: **{best_topic}**.\n\n- Punkt A\n- Punkt B", "closing": "Fragen? Melde dich bei HR.", "good_label": "Stark!"},
        "sections": [
            {"key": "custom:abc1", "enabled": True, "title": "Liebe/r {leader}", "body": "Gesamt: {overall} <script>alert(1)</script>\n\nZweiter Absatz"},
            {"key": "custom:bad key", "enabled": True, "title": "x", "body": "y"},  # ungueltige ID -> verworfen
            {"key": "guide", "enabled": True, "title": "Lesehilfe"},
        ],
    })
    assert [s["key"] for s in cfg["sections"] if s["key"].startswith("custom")] == ["custom:abc1"]
    assert cfg["sections"][0]["body"].startswith("Gesamt")
    html = report_render.render_html(report_render.SAMPLE, cfg)
    assert "Liebe/r Alex Beispiel" in html and "Gesamt: 3,8" in html.replace("3,80", "3,8") or "Gesamt: 3,8" in html
    assert "<script>" not in html and "&lt;script&gt;" in html  # Eingaben werden maskiert
    assert "<strong>Wertschätzung</strong>" in html and "<li>Punkt A</li>" in html  # Markdown-lite
    assert "Stark!" in html and "Fragen? Melde dich bei HR." in html
    assert html.index("Liebe/r") < html.index("Lesehilfe")


def test_text_override_equal_to_default_is_dropped_and_limits():
    d = report_layout.TEXTS["good_label"][1]
    cfg = report_layout.normalize({"texts": {"good_label": d, "unbekannt": "x"}, "sections": [{"key": f"custom:a{i}", "title": "t"} for i in range(15)]})
    assert cfg["texts"] == {}
    assert len([s for s in cfg["sections"] if s["key"].startswith("custom:")]) == report_layout.MAX_CUSTOM


def test_pptx_contains_custom_block_and_closing():
    from pptx import Presentation
    cfg = report_layout.normalize({"texts": {"closing": "Danke fürs Lesen"}, "sections": [{"key": "custom:x1", "title": "Eigener Text", "body": "- Eins\n- Zwei"}]})
    prs = Presentation(io.BytesIO(report_render.build_pptx(report_render.SAMPLE, cfg)))
    text = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    assert "Eigener Text" in text and "• Eins" in text and "Danke fürs Lesen" in text


def test_layout_api_exposes_text_catalog_and_placeholders(client, seeded_users):
    h = _h(client, "T-ADMIN")
    r = client.get("/report-layout", headers=h).json()
    assert {t["key"] for t in r["texts"]} >= {"guide_body", "good_label", "closing"}
    assert {p["key"] for p in r["placeholders"]} >= {"leader", "overall", "best_topic"}
    cfg = r["config"]
    cfg["texts"] = {"good_label": "Stark!"}
    cfg["sections"].append({"key": "custom:n1", "enabled": True, "title": "Neu", "body": "Text {leader}"})
    saved = client.put("/report-layout", json={"config": cfg}, headers=h).json()["config"]
    assert saved["texts"] == {"good_label": "Stark!"} and any(s["key"] == "custom:n1" for s in saved["sections"])
    assert "Text Alex Beispiel" in client.post("/report-layout/preview", json={"config": saved}, headers=h).json()["html"]
