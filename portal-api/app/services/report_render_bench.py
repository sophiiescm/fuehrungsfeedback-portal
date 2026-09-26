"""Benchmark-Exporte fuer Admins (nur Aggregate, unterdrueckte Gruppen entfallen)."""

from __future__ import annotations

import io


def benchmark_table(data: dict) -> tuple[list[str], list[list]]:
    """Aus der /benchmark-Antwort: Mittelwerte je Fachbereich und Dimension."""
    dims = sorted({k for g in data.values() if not g.get("suppressed") and isinstance(g.get("leaders"), list) for m in g["leaders"] for k in m["dimensions"]})
    rows = []
    for fb, g in sorted(data.items()):
        if fb == "_gesamt" or g.get("suppressed") or not isinstance(g.get("leaders"), list):
            continue
        ms = [m["dimensions"] for m in g["leaders"]]
        rows.append([fb, len(ms)] + [
            round(sum(m[d] for m in ms if d in m) / sum(1 for m in ms if d in m), 2) if any(d in m for m in ms) else ""
            for d in dims
        ])
    return dims, rows


def benchmark_xlsx(dims: list[str], rows: list[list], accent: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Benchmark"
    ws.append(["Fachbereich", "Führungskräfte"] + dims)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=accent.lstrip("#").upper())
    for r in rows:
        ws.append(r)
    ws.column_dimensions["A"].width = 24
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def benchmark_pptx(round_name: str, dims: list[str], rows: list[list], accent: str) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt

    acc = RGBColor.from_string(accent.lstrip("#").upper())
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = acc
    bg.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.5), Inches(2))
    tb.text_frame.text = f"Benchmark: {round_name}"
    p = tb.text_frame.paragraphs[0]
    p.font.size, p.font.bold, p.font.color.rgb = Pt(40), True, RGBColor(255, 255, 255)
    if rows and dims:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        t = s.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12), Inches(0.9))
        t.text_frame.text = "Themen nach Fachbereich (Ø, Skala 1–5)"
        p = t.text_frame.paragraphs[0]
        p.font.size, p.font.bold, p.font.color.rgb = Pt(28), True, acc
        cd = CategoryChartData()
        cd.categories = dims
        for r in rows:
            cd.add_series(str(r[0]), [v if v != "" else 0 for v in r[2:]])
        ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.3), Inches(12), Inches(5.6), cd).chart
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 1, 5
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
